"""
Clean the market price dataset and export every dataset used by KrishiSetu
into a single Excel workbook (dataset/KrishiSetu_Dataset.xlsx).

What it does:
  1. Removes exact duplicate rows from backend/cleaned_combine_wsrjc.csv
     (the CSV is what the price-forecasting API reads).
  2. Writes an Excel workbook where every row and column is visible:
       - Market Prices    : full price dataset (+ a real Date column)
       - Summary          : records / average price per state & commodity
       - Crop Nutrients   : NPK ranges + fertilizer per crop
       - Leaf Diseases    : disease classes with description & treatment
       - Data Dictionary  : meaning of every column
       - Data Quality     : what was cleaned and what to watch out for

Run from the project root:
    python backend/scripts/export_dataset.py
"""

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
CSV_PATH = ROOT / "backend" / "cleaned_combine_wsrjc.csv"
NPK_PATH = ROOT / "src" / "data" / "commodity_npk_fertilizer_data.json"
DISEASE_PATH = ROOT / "src" / "data" / "plant_disease_classes.json"
OUT_DIR = ROOT / "dataset"
OUT_PATH = OUT_DIR / "KrishiSetu_Dataset.xlsx"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F7A4D")
PRICE_COL = "Modal Price (Rs./Quintal)"
MSP_COL = "MSP (Rs./Quintal)"


def write_sheet(wb, title, df, widths=None, wrap_cols=()):
    """Write a DataFrame to a sheet with a styled, frozen, filterable header."""
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])

    for idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(idx)
        if widths and col in widths:
            width = widths[col]
        else:
            sample = df[col].astype(str).head(2000)
            width = min(max(len(str(col)), int(sample.str.len().max() or 0)) + 3, 60)
        ws.column_dimensions[letter].width = width
        ws.column_dimensions[letter].hidden = False
        if col in wrap_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{ws.max_row}"
    ws.row_dimensions[1].height = 32
    return ws


def remove_duplicate_lines(path):
    """Drop exact duplicate rows in place, keeping the file's original formatting."""
    lines = path.read_bytes().split(b"\n")
    seen, kept = set(), []
    for line in lines:
        if line and line in seen:
            continue
        seen.add(line)
        kept.append(line)
    path.write_bytes(b"\n".join(kept))
    return len(lines) - len(kept)


def main():
    removed_duplicates = remove_duplicate_lines(CSV_PATH)
    df = pd.read_csv(CSV_PATH)
    original_rows = len(df) + removed_duplicates

    df["Date"] = pd.to_datetime(df[["year", "month", "day"]]).dt.date
    df = df.sort_values(["State Name", "District Name", "Commodity", "Date", "Market Name"]).reset_index(drop=True)

    # Rows whose price is far away from the typical price of that commodity
    median = df.groupby("Commodity")[PRICE_COL].transform("median")
    suspicious = df[(df[PRICE_COL] > median * 4) | (df[PRICE_COL] < median / 4)]
    msp_zero = int((df[MSP_COL] == 0).sum())

    summary = (
        df.groupby(["State Name", "Commodity"])
        .agg(
            Districts=("District Name", "nunique"),
            Markets=("Market Name", "nunique"),
            Records=(PRICE_COL, "size"),
            **{
                "Avg Modal Price": (PRICE_COL, "mean"),
                "Min Modal Price": (PRICE_COL, "min"),
                "Max Modal Price": (PRICE_COL, "max"),
            },
            **{"First Date": ("Date", "min"), "Last Date": ("Date", "max")},
        )
        .reset_index()
    )
    summary["Avg Modal Price"] = summary["Avg Modal Price"].round(2)

    npk = pd.DataFrame(
        [
            {
                "Crop": crop,
                "Nitrogen N (kg/ha)": info["N_range_kg_per_ha"],
                "Phosphorus P (kg/ha)": info["P_range_kg_per_ha"],
                "Potassium K (kg/ha)": info["K_range_kg_per_ha"],
                "Preferred Fertilizer": info["preferred_fertilizer"],
            }
            for crop, info in json.loads(NPK_PATH.read_text(encoding="utf-8")).items()
        ]
    )

    diseases = pd.DataFrame(
        [
            {
                "Class": d["disease_name"].strip(),
                "Plant": d["disease_name"].strip().split("_")[0],
                "Description": d.get("description", ""),
                "Treatment": d.get("treatement", ""),
                "Fertilizer / Fungicide": d.get("fertilizerName", ""),
                "Watering": (d.get("waterLevel") or d.get("waterLevel ") or ""),
            }
            for d in json.loads(DISEASE_PATH.read_text(encoding="utf-8"))
        ]
    )

    dictionary = pd.DataFrame(
        [
            ("Market Prices", "State Name", "Indian state where the mandi (market) is located"),
            ("Market Prices", "District Name", "District of the mandi"),
            ("Market Prices", "Market Name", "Name of the APMC mandi"),
            ("Market Prices", "Commodity", "Crop traded: Cotton, Wheat, Rice, Jowar, Sugarcane"),
            ("Market Prices", MSP_COL, "Reference / minimum price recorded for the day (Rs. per quintal)"),
            ("Market Prices", PRICE_COL, "Most common trading price of the day (Rs. per quintal) - the value the model forecasts"),
            ("Market Prices", "day / month / year", "Date of the price record, split into parts (used by the model)"),
            ("Market Prices", "Date", "Same date as one Excel date column, for easy sorting / filtering"),
            ("Model features", "lag_1", "Last known modal price before the prediction date"),
            ("Model features", "lag_7", "Modal price 7 records before the prediction date"),
            ("Model features", "rolling_7", "Average of the last 7 modal prices"),
            ("Model features", "day_of_week / month / quarter", "Calendar features of the prediction date"),
            ("Model features", "State_enc / District_enc", "Label-encoded state and district"),
        ],
        columns=["Sheet / Group", "Column", "Meaning"],
    )

    quality = pd.DataFrame(
        [
            ("Rows in source CSV", original_rows),
            ("Exact duplicate rows removed", removed_duplicates),
            ("Rows kept", len(df)),
            ("Date range", f"{df['Date'].min()} to {df['Date'].max()}"),
            ("States", df["State Name"].nunique()),
            ("Districts", df["District Name"].nunique()),
            ("Markets", df["Market Name"].nunique()),
            ("Rows with MSP = 0 (kept, MSP is not used by the model)", msp_zero),
            ("Rows with price > 4x or < 1/4 of commodity median (kept, listed below)", len(suspicious)),
            ("Missing values", int(df.isna().sum().sum())),
        ],
        columns=["Check", "Value"],
    )

    OUT_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    write_sheet(wb, "Market Prices", df, widths={"Date": 13, "day": 7, "month": 8, "year": 8})
    for row in wb["Market Prices"].iter_rows(min_row=2, min_col=10, max_col=10):
        row[0].number_format = "dd-mmm-yyyy"
    write_sheet(wb, "Summary", summary)
    write_sheet(wb, "Crop Nutrients", npk)
    write_sheet(
        wb, "Leaf Diseases", diseases,
        widths={"Class": 34, "Plant": 16, "Description": 60, "Treatment": 60, "Fertilizer / Fungicide": 32, "Watering": 50},
        wrap_cols=("Description", "Treatment", "Watering"),
    )
    write_sheet(wb, "Data Dictionary", dictionary, widths={"Sheet / Group": 18, "Column": 30, "Meaning": 90})
    qws = write_sheet(wb, "Data Quality", quality, widths={"Check": 70, "Value": 28})
    qws.auto_filter.ref = None

    start = qws.max_row + 3
    qws.cell(row=start - 1, column=1, value="Suspicious price rows (check before trusting)").font = Font(bold=True)
    cols = ["State Name", "District Name", "Market Name", "Commodity", PRICE_COL, "Date"]
    for c, name in enumerate(cols, start=1):
        cell = qws.cell(row=start, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for r, row in enumerate(suspicious[cols].itertuples(index=False), start=start + 1):
        for c, v in enumerate(row, start=1):
            qws.cell(row=r, column=c, value=v)
    for c in range(3, 7):
        qws.column_dimensions[get_column_letter(c)].width = 22

    wb.save(OUT_PATH)
    print(f"Removed {removed_duplicates} duplicate rows, kept {len(df)}")
    print(f"Workbook written to {OUT_PATH}")


if __name__ == "__main__":
    main()
